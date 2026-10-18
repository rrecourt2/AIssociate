import json

import openpyxl  # dev-only: writes a realistic export with shared strings, like Excel does
import pytest

import fin_table


@pytest.fixture
def spread(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Spread"
    ws.append(["EUR m", 2023, 2024, 2025])          # numeric year headers, as many exports use
    ws.append(["Total revenue", 90, 100, 120])
    ws.append(["EBITDA", 15, 20, 18])
    ws.append(["Net debt", 40, 45, None])            # missing value
    ws.append(["Comment", "n/a", "x", "y"])          # text in a value row
    ws["E1"], ws["E2"] = "Example Group", "consolidated"
    path = tmp_path / "spread.xlsx"
    wb.save(path)
    return path


def mapping(**rows):
    base = {"sheet": "Spread", "label_column": "A", "header_row": 1, "unit": "EUR m",
            "identity": {"entity": {"value": "Example Group", "cell": "E1"},
                         "perimeter": {"value": "consolidated", "cell": "E2"},
                         "unit": {"value": "EUR m", "cell": "A1"}},
            "periods": ["2024", "2025"], "rows": [
                {"label": "Revenue", "source": "total revenue"},
                {"label": "EBITDA", "source": "EBITDA"},
                {"label": "EBITDA margin", "formula": "{EBITDA} / {Revenue}", "format": "pct"},
                {"label": "Revenue growth", "formula": "growth({Revenue})", "format": "pct"},
                {"label": "Net debt", "source": "Net debt"},
                {"label": "Net debt / EBITDA", "formula": "{Net debt} / {EBITDA}", "format": "x"}]}
    base.update(rows)
    return base


def test_table_values_sources_and_markdown(spread):
    t = fin_table.build_table(spread, mapping())
    rev = t["rows"][0]
    assert rev["values"] == [100.0, 120.0] and rev["source"] == ["Spread!C2", "Spread!D2"]
    md = fin_table.markdown(t)
    assert "| EBITDA margin | 20.0% | 15.0% |" in md
    assert "| Revenue growth | n/a | 20.0% |" in md       # first period has no prior in the table
    assert "| Net debt / EBITDA | 2.2x | n/a |" in md      # missing input stays n/a, never 0


def test_located_identity_is_retained(spread):
    m = mapping()
    table = fin_table.build_table(spread, m)
    assert table["identity"] == m["identity"]
    assert table["export"]["sheet"] == "Spread"


def test_reviewed_external_identity_is_retained(spread):
    wb = openpyxl.load_workbook(spread)
    for ref in ("A1", "E1", "E2"):
        wb["Spread"][ref] = None
    wb.save(spread)
    identity = {key: {"value": value, "basis": "Synthetic accounts, cover and note 1",
                      "reviewed_by": "Synthetic officer"}
                for key, value in (("entity", "Example Group"), ("perimeter", "consolidated"),
                                   ("unit", "EUR m"))}
    table = fin_table.build_table(spread, mapping(identity=identity))
    assert table["identity"] == identity
    assert table["rows"][0]["values"] == [100.0, 120.0]


@pytest.mark.parametrize("missing", ["identity", "entity", "perimeter", "unit"])
def test_missing_identity_is_refused(spread, missing):
    m = mapping()
    del (m if missing == "identity" else m["identity"])[missing]
    with pytest.raises(fin_table.TableError, match="identity"):
        fin_table.build_table(spread, m)


@pytest.mark.parametrize("field, value", [
    ("entity", "Other Group"), ("perimeter", "standalone"), ("unit", "USD m"),
    ("unit", "EUR thousands"),
])
def test_wrong_identity_is_refused(spread, field, value):
    m = mapping()
    m["identity"][field]["value"] = value
    with pytest.raises(fin_table.TableError, match=f"identity.*{field}"):
        fin_table.build_table(spread, m)


@pytest.mark.parametrize("entry", [
    None, "Example Group", {}, {"value": "Example Group"},
    {"value": None, "cell": "E1"}, {"value": " ", "cell": "E1"},
    {"value": 2024, "cell": "C1"}, {"value": "Example Group", "cell": "Other!E1"},
    {"value": "Example Group", "cell": "E0"}, {"value": "Example Group", "cell": ""},
    {"value": "Example Group", "cell": "E1", "basis": "override", "reviewed_by": "Officer"},
    {"value": "Example Group", "cell": "E1", "unexpected": "ignored"},
    {"value": "Example Group", "basis": " ", "reviewed_by": "Officer"},
    {"value": "Example Group", "basis": "Cover page", "reviewed_by": ""},
    {"value": "Example Group", "basis": "Cover page", "reviewed_by": True},
])
def test_malformed_identity_entry_is_refused(spread, entry):
    m = mapping()
    m["identity"]["entity"] = entry
    with pytest.raises(fin_table.TableError, match="identity.*entity"):
        fin_table.build_table(spread, m)


@pytest.mark.parametrize("identity", [None, [], {}, {"entity": {}, "perimeter": {}, "unit": {}, "other": {}}])
def test_malformed_identity_object_is_refused(spread, identity):
    with pytest.raises(fin_table.TableError, match="identity"):
        fin_table.build_table(spread, mapping(identity=identity))


@pytest.mark.parametrize("actual", [None, " ", 2024, False])
def test_identity_cell_requires_nonempty_text(spread, actual):
    wb = openpyxl.load_workbook(spread)
    wb["Spread"]["E1"] = actual
    wb.save(spread)
    m = mapping()
    m["identity"]["entity"]["value"] = str(actual)
    with pytest.raises(fin_table.TableError, match="identity.*entity"):
        fin_table.build_table(spread, m)


def test_identity_cells_are_checked_in_selected_sheet(spread):
    wb = openpyxl.load_workbook(spread)
    other = wb.copy_worksheet(wb["Spread"])
    other.title = "Other"
    other["E1"] = "Other Group"
    wb.save(spread)
    with pytest.raises(fin_table.TableError, match="identity.*entity"):
        fin_table.build_table(spread, mapping(sheet="Other"))


@pytest.mark.parametrize("unit", [None, "", "USD m", "EUR thousands"])
def test_display_unit_must_agree_with_identity(spread, unit):
    with pytest.raises(fin_table.TableError, match="unit"):
        fin_table.build_table(spread, mapping(unit=unit))


def test_missing_display_unit_is_refused(spread):
    m = mapping()
    del m["unit"]
    with pytest.raises(fin_table.TableError, match="unit"):
        fin_table.build_table(spread, m)


@pytest.mark.parametrize("change, message", [
    ({"periods": ["2024", "2030"]}, "period '2030'"),
    ({"rows": [{"label": "X", "source": "Turnover"}]}, "row 'Turnover'"),
    ({"rows": [{"label": "X", "formula": "{A} * {B}"}]}, "unsupported formula"),
    ({"rows": [{"label": "X", "formula": "{Revenue} / {EBITDA}"}]}, "not an earlier row"),
    ({"rows": [{"label": "X", "source": "Comment"}]}, "not numbers"),
    ({"sheet": "Other"}, "no sheet"),
])
def test_refuses_instead_of_guessing(spread, change, message):
    with pytest.raises(fin_table.TableError, match=message):
        fin_table.build_table(spread, mapping(**change))


def test_growth_on_negative_base_is_not_reported():
    assert fin_table.compute("growth({A})", {"A": [-10.0, 5.0, 10.0]}) == [None, None, 1.0]


def test_cli_writes_json_and_markdown(spread, tmp_path, capsys):
    m = tmp_path / "mapping.json"
    m.write_text(json.dumps(mapping()))
    out = tmp_path / "table.json"
    assert fin_table.main([str(spread), str(m), "--json", str(out)]) == 0
    assert json.loads(out.read_text())["rows"][2]["formula"] == "{EBITDA} / {Revenue}"
    assert "| EBITDA margin | 20.0% | 15.0% |" in capsys.readouterr().out


def test_ambiguous_rows_and_periods_are_refused(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Spread"
    ws.append(["EUR m", 2024, 2025, 2025])
    ws.append(["EBITDA", 20, 18, 17])
    ws.append(["EBITDA", 1, 2, 3])
    ws["E1"], ws["E2"] = "Example Group", "consolidated"
    path = tmp_path / "dup.xlsx"
    wb.save(path)
    m = mapping(periods=["2024"], rows=[{"label": "EBITDA", "source": "EBITDA"}])
    with pytest.raises(fin_table.TableError, match="ambiguous.*row 'EBITDA'"):
        fin_table.build_table(path, m)
    with pytest.raises(fin_table.TableError, match="ambiguous.*period '2025'"):
        fin_table.build_table(path, dict(m, periods=["2025"], rows=[]))


def test_formula_cells_need_an_explicitly_trusted_recalculated_export(tmp_path):
    path, m = formula_export(tmp_path)
    with pytest.raises(fin_table.TableError, match="formulas"):
        fin_table.build_table(path, m)
    assert fin_table.build_table(path, dict(m, trust_formula_values=True, recalculation_basis="Officer: recalculated in Excel and saved 2026-10-02"))["rows"][0]["values"] == [30.0]


def test_expected_cells_identify_the_export(spread):
    m = mapping(expect={"A1": "EUR m", "A2": "Total revenue"})
    assert fin_table.build_table(spread, m)["rows"][0]["values"] == [100.0, 120.0]
    with pytest.raises(fin_table.TableError, match="not the expected export"):
        fin_table.build_table(spread, mapping(expect={"A1": "USD m"}))


def fin_table_inject_formula(path, ref, formula):
    """Add a formula to a cell that keeps its cached value, as Excel saves it."""
    import re
    import zipfile
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    sheet, count = re.subn(rf'(<c r="{ref}"[^>]*>)', rf"\1<f>{formula}</f>",
                           parts["xl/worksheets/sheet1.xml"].decode(), count=1)
    assert count == 1
    parts["xl/worksheets/sheet1.xml"] = sheet.encode()
    with zipfile.ZipFile(path, "w") as z:
        for n, d in parts.items():
            z.writestr(n, d)


def formula_export(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Spread'
    for row in (['EUR m', '2025'], ['Revenue', 120], ['Gross profit', 30],
                ['Entity', 'Example Group'], ['Perimeter', 'consolidated']):
        ws.append(row)
    path = tmp_path / 'formula.xlsx'
    wb.save(path)
    fin_table_inject_formula(path, 'B3', 'B2*0.25')
    m = mapping(periods=['2025'], rows=[{'label': 'Gross profit', 'source': 'Gross profit'}])
    m['identity']['entity']['cell'] = 'B4'
    m['identity']['perimeter']['cell'] = 'B5'
    return path, m


@pytest.mark.parametrize('trust, basis', [(True, ''), (True, '  '), ('yes', 'recalculated')])
def test_formula_trust_requires_boolean_opt_in_and_recalculation_basis(tmp_path, trust, basis):
    path, m = formula_export(tmp_path)
    with pytest.raises(fin_table.TableError, match='trust_formula_values|recalculation_basis'):
        fin_table.build_table(path, dict(m, trust_formula_values=trust, recalculation_basis=basis))


def test_export_and_trusted_formula_provenance_are_retained(tmp_path):
    import hashlib
    path, m = formula_export(tmp_path)
    table = fin_table.build_table(path, dict(m, trust_formula_values=True, recalculation_basis='Excel: full recalculation and save, officer 2026-10-02'))
    assert table['export']['path'] == str(path.resolve())
    assert table['export']['sha256'] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert table['export']['formula_cells'] == ['B3']
    assert table['export']['recalculation_basis'].startswith('Excel:')
    assert table['export']['trust_formula_values'] is True


@pytest.mark.parametrize('kind', ['periods', 'labels', 'both_source_and_formula'])
def test_ambiguous_mapping_is_refused(spread, kind):
    if kind == 'periods':
        m = mapping(periods=['2024', '2024'])
    elif kind == 'labels':
        m = mapping(rows=[{'label': 'Revenue', 'source': 'Total revenue'}, {'label': ' revenue ', 'source': 'EBITDA'}])
    else:
        m = mapping(rows=[{'label': 'Revenue', 'source': 'Total revenue', 'formula': 'growth({Revenue})'}])
    with pytest.raises(fin_table.TableError, match='duplicate|exactly one'):
        fin_table.build_table(spread, m)


@pytest.mark.parametrize('target', ['spread', 'mapping', 'existing', 'symlink'])
def test_json_output_cannot_overwrite_inputs_or_existing_files(spread, tmp_path, target, capsys):
    m = tmp_path / 'mapping.json'
    m.write_text(json.dumps(mapping()))
    out = tmp_path / 'table.json'
    if target == 'spread':
        out = spread
    elif target == 'mapping':
        out = m
    elif target == 'existing':
        out.write_text('prior output')
    else:
        out.symlink_to(spread)
    before = {p: p.read_bytes() for p in (spread, m, out)}
    assert fin_table.main([str(spread), str(m), '--json', str(out)]) == 3
    assert 'REFUSED' in capsys.readouterr().err
    assert all(p.read_bytes() == data for p, data in before.items())


def test_atomic_output_failure_leaves_no_partial_json(spread, tmp_path, monkeypatch):
    import os
    m = tmp_path / 'mapping.json'
    m.write_text(json.dumps(mapping()))
    out = tmp_path / 'table.json'
    monkeypatch.setattr(os, 'link', lambda *a, **kw: (_ for _ in ()).throw(OSError('publish failed')))
    assert fin_table.main([str(spread), str(m), '--json', str(out)]) == 3
    assert not out.exists()
    assert not list(tmp_path.glob('*.tmp'))
